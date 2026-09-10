"""The Coca-Cola personnel roster: Coca-Cola card number -> person's name.

Coca-Cola export a card list (`CardNo`, `User`, `FirstName`, `LastName`) where
CardNo is the `DDD-DDDDD` form of the card. Our kiosk only ever sees the POS
id, so a scan is joined to a person by converting POS -> Coca-Cola code
(see cardcode.pos_to_cc) and looking the result up here.

Why a separate table instead of typing names onto cards by hand: the roster is
866 people, and hand-typing is how you get "ურა ლაბაშვილი" where the official
list says "უშანგი ლაზაშვილი". Importing keeps our reports spelled exactly the
way Coca-Cola's own records spell them.

The roster is reference data, NOT access control: an unknown card still scans
and still eats (see scan_service). It only decides which name a report shows.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

from sqlmodel import Session, select

from .cardcode import normalize_cc_code
from .models import RosterEntry

# Header cells we recognise, lower-cased. The Coca-Cola export uses the first
# four; the others are tolerated so a re-exported/edited file still imports.
_H_CARD = {"cardno", "card no", "card_no", "code", "კოდი", "ბარათი"}
_H_FIRST = {"firstname", "first name", "first_name", "სახელი"}
_H_LAST = {"lastname", "last name", "last_name", "გვარი"}
_H_USER = {"user", "name", "fullname", "full name", "სრული სახელი"}


@dataclass
class RosterReport:
    added: int = 0
    updated: int = 0
    invalid: list[tuple[int, str]] = field(default_factory=list)  # (row, reason)
    total_rows: int = 0

    def as_dict(self) -> dict:
        return {
            "added": self.added,
            "updated": self.updated,
            "invalid_count": len(self.invalid),
            "total_rows": self.total_rows,
            "invalid": [{"row": r, "reason": why} for r, why in self.invalid],
        }


def _cell_text(value) -> str:  # noqa: ANN001
    """Coerce a spreadsheet cell to clean text (floats -> ints, no '.0')."""
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, int):
        return str(value)
    return str(value).lstrip("﻿").strip()


class RosterFormatError(RuntimeError):
    """A file we cannot read — surfaced to the operator, not a 500."""


def _rows_from_xls(data: bytes) -> list[list[str]]:
    """Read a legacy .xls (BIFF). The Coca-Cola export is this format."""
    try:
        import xlrd
    except ImportError:
        # xlrd arrives with a normal update, but quick-start.bat launches
        # without running pip, so a kiosk can be on new code with old deps.
        # Say what to do instead of dying with ModuleNotFoundError.
        raise RosterFormatError(
            "ძველი .xls ფორმატის წასაკითხად საჭიროა დამატებითი კომპონენტი. "
            "გაუშვით start.bat კიოსკზე (ან შეინახეთ ფაილი .xlsx ფორმატში "
            "და თავიდან ატვირთეთ)."
        )

    book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    return [[_cell_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)]


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = book.active
    out = [[_cell_text(v) for v in row]
           for row in sheet.iter_rows(values_only=True)]
    book.close()
    return out


def _rows_from_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    return [[_cell_text(c) for c in row] for row in csv.reader(io.StringIO(text))]


def parse_rows(filename: str, data: bytes) -> list[list[str]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _rows_from_csv(data)
    if name.endswith(".xls"):
        return _rows_from_xls(data)
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _rows_from_xlsx(data)
    # Unknown extension: sniff the OLE2 magic of a legacy .xls, else try xlsx.
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _rows_from_xls(data)
    try:
        return _rows_from_xlsx(data)
    except Exception:  # noqa: BLE001
        return _rows_from_csv(data)


def _find_columns(header: list[str]) -> dict[str, int]:
    """Locate the columns by header name; fall back to the export's order."""
    found: dict[str, int] = {}
    for idx, cell in enumerate(header):
        key = cell.strip().lower()
        if key in _H_CARD and "card" not in found:
            found["card"] = idx
        elif key in _H_FIRST and "first" not in found:
            found["first"] = idx
        elif key in _H_LAST and "last" not in found:
            found["last"] = idx
        elif key in _H_USER and "user" not in found:
            found["user"] = idx
    if "card" not in found:
        # No recognisable header — assume the Coca-Cola layout:
        # CardNo | User | FirstName | LastName
        found = {"card": 0, "user": 1, "first": 2, "last": 3}
    return found


def _full_name(row: list[str], cols: dict[str, int]) -> str:
    """Prefer 'FirstName LastName'; fall back to the combined 'User' column."""
    def at(key: str) -> str:
        idx = cols.get(key)
        return row[idx].strip() if idx is not None and idx < len(row) else ""

    first, last = at("first"), at("last")
    if first or last:
        return " ".join(p for p in (first, last) if p)
    return at("user")


def import_roster(session: Session, filename: str, data: bytes) -> RosterReport:
    """Load/refresh the roster. Existing codes are UPDATED, not duplicated."""
    rows = parse_rows(filename, data)
    report = RosterReport()
    if not rows:
        return report

    cols = _find_columns(rows[0])
    # Skip row 1 only when it really is a header (its card cell is not a code).
    start = 1 if not normalize_cc_code(rows[0][cols.get("card", 0)]
                                       if cols.get("card", 0) < len(rows[0]) else "") else 0

    existing = {e.cc_code: e for e in session.exec(select(RosterEntry)).all()}
    seen: set[str] = set()

    for idx, row in enumerate(rows[start:], start=start + 1):
        card_idx = cols.get("card", 0)
        raw = row[card_idx] if card_idx < len(row) else ""
        if not raw.strip():
            continue                      # blank line, not an error
        report.total_rows += 1

        code = normalize_cc_code(raw)
        if not code:
            report.invalid.append((idx, f"არასწორი კოდი: {raw[:24]}"))
            continue
        if code in seen:
            continue                      # duplicate inside the file
        seen.add(code)

        name = _full_name(row, cols)
        if not name:
            report.invalid.append((idx, f"სახელი ცარიელია ({code})"))
            continue

        entry = existing.get(code)
        if entry is None:
            session.add(RosterEntry(cc_code=code, full_name=name))
            report.added += 1
        elif entry.full_name != name:
            entry.full_name = name
            session.add(entry)
            report.updated += 1

    if report.added or report.updated:
        session.commit()
    return report


def upsert_entry(session: Session, cc_code: str, full_name: str) -> bool:
    """Add/refresh ONE roster name. Returns True when a new entry was created.

    The bulk import is for Coca-Cola's whole export; this is the single-person
    path behind the admin page's add box, used when somebody is identified by
    their DDD-DDDDD code rather than by a card that has physically tapped.
    """
    code = normalize_cc_code(cc_code)
    name = (full_name or "").strip()
    if not code or not name:
        return False
    entry = session.exec(
        select(RosterEntry).where(RosterEntry.cc_code == code)
    ).first()
    if entry is None:
        session.add(RosterEntry(cc_code=code, full_name=name))
        session.commit()
        return True
    if entry.full_name != name:
        entry.full_name = name
        session.add(entry)
        session.commit()
    return False


def names_by_cc_code(session: Session) -> dict[str, str]:
    """Whole roster as {cc_code: full_name} — one query for bulk reporting."""
    return {e.cc_code: e.full_name for e in session.exec(select(RosterEntry)).all()}


def lookup_name(session: Session, cc_code: str) -> str:
    if not cc_code:
        return ""
    entry = session.exec(
        select(RosterEntry).where(RosterEntry.cc_code == cc_code)
    ).first()
    return entry.full_name if entry else ""
