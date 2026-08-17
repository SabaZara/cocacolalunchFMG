"""The tap log: every card tap, allowed and denied.

`scans` records granted meals only, so before this a denied tap left no trace:
somebody turned away at the reader was invisible afterwards, and a quiet day
could not be told apart from a day full of rejections.
"""
from __future__ import annotations

import io
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


def _login(ctx):
    ctx["client"].post("/api/login", headers=ctx["headers"],
                       json={"username": ctx["admin_user"],
                             "password": ctx["admin_pass"]})


def _cap(ctx, card_id, limit=1):
    """Give a card a real limit. New cards are UNLIMITED by default, so a test
    that needs a limit-denial must ask for one explicitly."""
    c, H = ctx["client"], ctx["headers"]
    c.post("/api/scan", json={"card_id": card_id})          # registers it
    pid = c.get(f"/api/people?q={card_id}", headers=H).json()[0]["id"]
    c.put(f"/api/people/{pid}", headers=H, json={"daily_limit": limit})
    # clear the registering meal so the test starts from zero
    c.post(f"/api/people/{pid}/ate", headers=H, json={"ate": False})


def _roster_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["CardNo", "User", "FirstName", "LastName"])
    for code, first, last in rows:
        ws.append([code, f"{first} {last}", first, last])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_logs_allowed_and_denied_taps(app_ctx):
    ctx = app_ctx
    c = ctx["client"]
    H = ctx["headers"]
    _login(ctx)

    _cap(ctx, "3377269862", 1)
    c.post("/api/scan", json={"card_id": "3377269862"})   # allowed
    c.post("/api/scan", json={"card_id": "3377269862"})   # denied: limit
    c.post("/api/scan", json={"card_id": "   "})          # denied: empty read

    d = date.today().isoformat()
    log = c.get(f"/api/reports/taplog?from={d}&to={d}", headers=H).json()

    # 1 registering tap from _cap + 3 taps here
    assert log["total"] == 4
    assert log["allowed"] == 2      # the registering tap + the first real one
    assert log["denied"] == 2
    reasons = {r["reason"]: r["count"] for r in log["by_reason"]}
    assert reasons["დღის ლიმიტი ამოიწურა"] == 1
    assert reasons["ბარათი ვერ წაიკითხა"] == 1
    # newest first
    assert log["rows"][0]["status"] == "DENIED"


def test_log_carries_name_and_code(app_ctx):
    """A denied tap must say WHO was turned away, not just which number."""
    ctx = app_ctx
    c = ctx["client"]
    H = ctx["headers"]
    _login(ctx)
    c.post("/api/people/roster-import", headers=H,
           files={"file": ("r.xlsx", _roster_xlsx([("077-03174", "ია", "ცუცქირიძე")]),
                           "application/octet-stream")})
    _cap(ctx, "3377269862", 1)
    c.post("/api/scan", json={"card_id": "3377269862"})
    c.post("/api/scan", json={"card_id": "3377269862"})   # denied

    d = date.today().isoformat()
    rows = c.get(f"/api/reports/taplog?from={d}&to={d}", headers=H).json()["rows"]
    denied = [r for r in rows if r["status"] == "DENIED"][0]
    assert denied["full_name"] == "ია ცუცქირიძე"
    assert denied["cc_code"] == "077-03174"
    assert denied["reason"] == "დღის ლიმიტი ამოიწურა"
    assert denied["status_ka"] == "უარყოფილი"


def test_status_filter_keeps_totals_whole(app_ctx):
    """Filtering the list must not change the counters above it."""
    ctx = app_ctx
    c = ctx["client"]
    H = ctx["headers"]
    _login(ctx)
    _cap(ctx, "1111111111", 1)
    c.post("/api/scan", json={"card_id": "1111111111"})
    c.post("/api/scan", json={"card_id": "1111111111"})   # denied

    d = date.today().isoformat()
    only_denied = c.get(f"/api/reports/taplog?from={d}&to={d}&status=DENIED",
                        headers=H).json()
    assert len(only_denied["rows"]) == 1
    assert only_denied["rows"][0]["status"] == "DENIED"
    # counters still describe the whole range (incl. _cap's registering tap)
    assert only_denied["total"] == 3
    assert only_denied["allowed"] == 2

    assert c.get(f"/api/reports/taplog?from={d}&to={d}&status=NOPE",
                 headers=H).status_code == 422


def test_logging_never_changes_meal_counting(app_ctx):
    """Meals are decided by `scans`; the log must stay purely observational."""
    ctx = app_ctx
    c = ctx["client"]
    H = ctx["headers"]
    _login(ctx)
    _cap(ctx, "2222222222", 1)
    for _ in range(5):
        c.post("/api/scan", json={"card_id": "2222222222"})

    d = date.today().isoformat()
    day = c.get(f"/api/reports/day?date={d}", headers=H).json()
    assert day["meals"] == 1            # limit 1 -> exactly one meal
    log = c.get(f"/api/reports/taplog?from={d}&to={d}", headers=H).json()
    # every attempt recorded: _cap's registering tap + the 5 here
    assert log["total"] == 6


def test_log_export_both_formats(app_ctx):
    ctx = app_ctx
    c = ctx["client"]
    H = ctx["headers"]
    _login(ctx)
    c.post("/api/people/roster-import", headers=H,
           files={"file": ("r.xlsx", _roster_xlsx([("077-03174", "ია", "ცუცქირიძე")]),
                           "application/octet-stream")})
    _cap(ctx, "3377269862", 1)
    c.post("/api/scan", json={"card_id": "3377269862"})
    c.post("/api/scan", json={"card_id": "3377269862"})

    d = date.today().isoformat()
    csv_body = c.get(f"/api/reports/taplog-export?from={d}&to={d}&format=csv",
                     headers=H).content.decode("utf-8-sig")
    assert "ია ცუცქირიძე" in csv_body
    assert "უარყოფილი" in csv_body and "ნებადართული" in csv_body
    assert "3377269862" in csv_body

    # Coca-Cola-only variant drops the POS id
    cc_only = c.get(f"/api/reports/taplog-export?from={d}&to={d}&format=csv&pos=0",
                    headers=H).content.decode("utf-8-sig")
    assert "077-03174" in cc_only
    assert "3377269862" not in cc_only

    # xlsx carries a summary sheet alongside the flat log
    x = c.get(f"/api/reports/taplog-export?from={d}&to={d}&format=xlsx", headers=H)
    wb = load_workbook(io.BytesIO(x.content))
    assert "ლოგი" in wb.sheetnames and "ჯამი" in wb.sheetnames
    summary = [[c.value for c in row] for row in wb["ჯამი"].iter_rows()]
    flat = {r[0]: r[1] for r in summary if r and r[0]}
    assert flat["სულ"] == 3          # incl. _cap's registering tap


def test_log_is_included_in_backups(app_ctx):
    """The log must survive in backups like every other table."""
    ctx = app_ctx
    c = ctx["client"]
    _login(ctx)
    _cap(ctx, "3377269862", 1)
    c.post("/api/scan", json={"card_id": "3377269862"})
    c.post("/api/scan", json={"card_id": "3377269862"})   # denied

    path = ctx["backup"].create_backup(reason="test")

    import sqlite3
    con = sqlite3.connect(path)
    try:
        tables = {r[0] for r in con.execute(
            "SELECT name FROM sqlite_master WHERE type='table'")}
        assert "tap_log" in tables, "tap log missing from the backup"
        n = con.execute("SELECT COUNT(*) FROM tap_log").fetchone()[0]
        assert n == 3               # incl. _cap's registering tap
        statuses = {r[0] for r in con.execute("SELECT status FROM tap_log")}
        assert statuses == {"ALLOWED", "DENIED"}
    finally:
        con.close()


def test_a_logging_failure_never_blocks_a_meal(app_ctx, monkeypatch):
    """The reader must keep working even if the log cannot be written."""
    ctx = app_ctx
    c = ctx["client"]
    svc = ctx["scan_service"]

    def boom(*a, **k):
        raise RuntimeError("disk full")

    monkeypatch.setattr(svc, "_log_tap", boom)
    r = c.post("/api/scan", json={"card_id": "3377269862"})
    assert r.status_code == 200
    assert r.json()["status"] == "ALLOWED", "a logging fault denied a meal"


def test_logs_page_is_served_and_gated(app_ctx):
    """The log lives on its own page, protected like admin and reports."""
    ctx = app_ctx
    c = ctx["client"]

    # No tunnel secret -> blocked, same as the other operator pages.
    assert c.get("/logs").status_code == 403

    r = c.get("/logs", headers=ctx["headers"])
    assert r.status_code == 200
    body = r.text
    assert "logs.js" in body
    assert "ბარათების ლოგი" in body
    # and it links back to the other two tabs
    assert '/admin' in body and '/reports' in body
